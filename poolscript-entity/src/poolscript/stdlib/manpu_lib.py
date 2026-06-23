"""
Módulo `manpu` da PoolScript — manipulação de arquivos.

Suporta: CSV, XLSX, XML, HTML, e arquivos de texto/código.

Uso:
    import manpu as mp

    ler = mp.read("arquivo.csv")
    post(ler)

    mp.write(content="Kleber", column=3, celula=12, target="arquivo.csv")
    mp.remove(value="Texto", amount="full", target="arquivo.html")
"""
from __future__ import annotations
import csv
import json
import os
from pathlib import Path
from typing import Any


def _ext(filepath: str) -> str:
    return Path(filepath).suffix.lower().lstrip(".")


class ManpuResult:
    """Resultado de operação da manpu."""
    def __init__(self, success: bool, status: str = ""):
        self._success = success
        self.status = status or ("Success" if success else "Error")

    def __eq__(self, other):
        if other is True:
            return self._success
        if other is False:
            return not self._success
        return self.status == other

    def __bool__(self):
        return self._success

    def __repr__(self):
        return self.status


def read(filepath: str) -> Any:
    """Lê um arquivo e retorna seu conteúdo formatado."""
    path = Path(filepath)
    if not path.is_file():
        return f"Error: arquivo não encontrado: {filepath}"

    ext = _ext(filepath)

    # CSV → lista de dicts
    if ext == "csv":
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            return [dict(row) for row in reader]

    # XLSX/EXCEL → lista de dicts
    if ext in ("xlsx", "xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
            return [dict(zip(headers, row)) for row in rows[1:]]
        except ImportError:
            return "Error: instale openpyxl para ler arquivos xlsx"

    # XML → dict
    if ext == "xml":
        try:
            import xml.etree.ElementTree as ET
            tree = ET.parse(path)
            root = tree.getroot()

            def _parse(el):
                result = {"tag": el.tag, "attrs": dict(el.attrib), "text": (el.text or "").strip(), "children": []}
                for child in el:
                    result["children"].append(_parse(child))
                return result

            return _parse(root)
        except Exception as e:
            return f"Error: {e}"

    # HTML → texto limpo
    if ext in ("html", "htm"):
        import re
        content = path.read_text(encoding="utf-8")
        clean = re.sub(r"<[^>]+>", "", content)
        clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
        return clean

    # JSON
    if ext == "json":
        with open(path, encoding="utf-8") as f:
            return json.load(f)

    # Texto puro — txt, py, js, ps, md, etc
    return path.read_text(encoding="utf-8")


def load(filepath: str) -> Any:
    """Carrega o arquivo pra uso em outros contextos (ex: corpo de email)."""
    path = Path(filepath)
    if not path.is_file():
        return f"Error: arquivo não encontrado: {filepath}"
    return path.read_bytes()


def write(content: Any = "", column: int = 0, celula: int = 0,
          target: str = "") -> ManpuResult:
    """
    Escreve conteúdo em um arquivo.
    Em CSV/XLSX: usa column e celula como índices.
    Em outros: adiciona ao final do arquivo.
    """
    path = Path(target)
    ext = _ext(target)

    try:
        if ext == "csv":
            rows = []
            if path.is_file():
                with open(path, newline="", encoding="utf-8") as f:
                    rows = list(csv.reader(f))

            # Expande se necessário
            while len(rows) <= celula:
                rows.append([])
            while len(rows[celula]) <= column:
                rows[celula].append("")

            rows[celula][column] = str(content)

            with open(path, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(rows)

        elif ext in ("xlsx", "xls"):
            try:
                import openpyxl
                if path.is_file():
                    wb = openpyxl.load_workbook(path)
                else:
                    wb = openpyxl.Workbook()
                ws = wb.active
                ws.cell(row=celula + 1, column=column + 1, value=content)
                wb.save(path)
            except ImportError:
                return ManpuResult(False, "Error: instale openpyxl")

        else:
            # Texto puro — adiciona ao final
            with open(path, "a", encoding="utf-8") as f:
                f.write(str(content))

        return ManpuResult(True)

    except Exception as e:
        return ManpuResult(False, f"Error: {e}")


def remove(value: str = "", content: str = "", column: int = 0,
           celula: int = 0, amount: str = "full",
           target: str = "") -> ManpuResult:
    """
    Remove conteúdo de um arquivo.
    amount='full' → remove tudo que corresponde
    amount='mei'  → remove metade do texto encontrado
    """
    path = Path(target)
    if not path.is_file():
        return ManpuResult(False, f"Error: arquivo não encontrado: {target}")

    ext = _ext(target)
    target_value = value or content

    try:
        if ext == "csv":
            rows = []
            with open(path, newline="", encoding="utf-8") as f:
                rows = list(csv.reader(f))

            if celula < len(rows) and column < len(rows[celula]):
                if amount == "full":
                    rows[celula][column] = ""
                elif amount == "mei":
                    val = rows[celula][column]
                    rows[celula][column] = val[:len(val)//2]

            with open(path, "w", newline="", encoding="utf-8") as f:
                csv.writer(f).writerows(rows)

        else:
            # Texto/HTML/outros
            text = path.read_text(encoding="utf-8")
            if not target_value:
                return ManpuResult(False, "Error: forneça value para remover")

            if amount == "full":
                text = text.replace(target_value, "")
            elif amount == "mei":
                idx = text.find(target_value)
                if idx != -1:
                    half = len(target_value) // 2
                    text = text[:idx] + target_value[half:] + text[idx + len(target_value):]

            path.write_text(text, encoding="utf-8")

        return ManpuResult(True)

    except Exception as e:
        return ManpuResult(False, f"Error: {e}")


class ManpuFile:
    """
    Arquivo aberto pelo mp.open() — mantém o arquivo aberto durante operações.
    Funciona com `using mp.open(...) as arq { ... }` da PoolScript.
    """

    def __init__(self, filepath: str):
        self.filepath = str(filepath)
        self._ext = _ext(self.filepath)
        self._path = Path(self.filepath)
        self._wb = None  # workbook xlsx
        self._rows = None  # rows csv
        self._text = None  # texto puro

        # Carrega o arquivo
        if self._ext == "csv":
            if self._path.is_file():
                with open(self._path, newline="", encoding="utf-8") as f:
                    self._rows = list(csv.reader(f))
            else:
                self._rows = []
        elif self._ext in ("xlsx", "xls"):
            try:
                import openpyxl
                if self._path.is_file():
                    self._wb = openpyxl.load_workbook(self._path)
                else:
                    import openpyxl as _xl
                    self._wb = _xl.Workbook()
            except ImportError:
                pass
        else:
            self._text = self._path.read_text(encoding="utf-8") if self._path.is_file() else ""

    def write(self, content: Any = "", column: int = 0,
              cell: Any = None, celula: Any = None,
              init: int = 0, sep: str = "\n",
              size: int = 0) -> ManpuResult:
        """
        Escreve conteúdo no arquivo aberto.

        cell/celula: índice da célula, ou 'full' pra preencher descendo linha a linha
        init: índice inicial do conteúdo (default 0)
        sep: separador pra dividir texto em células (default quebra de linha)
        size: tamanho fixo por célula (alternativa ao sep)
        column: coluna alvo (int ou 'full' pra todas)
        """
        try:
            cell_val = cell if cell is not None else celula

            # Divide o conteúdo em partes
            if isinstance(content, (bytes, bytearray)):
                text = content.decode("utf-8", errors="replace")
            else:
                text = str(content)

            if size > 0:
                parts = [text[i:i+size] for i in range(0, len(text), size)]
            else:
                parts = text.split(sep)

            parts = parts[init:]

            if self._ext == "csv":
                while len(self._rows) <= (init + len(parts)):
                    self._rows.append([])

                if cell_val == "full":
                    for i, part in enumerate(parts):
                        row_idx = i
                        while len(self._rows) <= row_idx:
                            self._rows.append([])
                        while len(self._rows[row_idx]) <= column:
                            self._rows[row_idx].append("")
                        self._rows[row_idx][column] = part
                else:
                    row_idx = int(cell_val)
                    while len(self._rows) <= row_idx:
                        self._rows.append([])
                    while len(self._rows[row_idx]) <= column:
                        self._rows[row_idx].append("")
                    self._rows[row_idx][column] = text

            elif self._ext in ("xlsx", "xls") and self._wb:
                ws = self._wb.active
                max_row = ws.max_row or 0

                if cell_val == "full":
                    if column == "full":
                        # preenche todas as colunas disponíveis
                        max_col = ws.max_column or 1
                        col_idx = 1
                        for part in parts:
                            if col_idx > max_col:
                                return ManpuResult(False, f"Error: Arquivo xlsx tem colunas insuficientes")
                            ws.cell(row=max_row + 1, column=col_idx, value=part)
                            col_idx += 1
                    else:
                        for i, part in enumerate(parts):
                            row = max_row + 1 + i
                            if ws.max_row and row > ws.max_row + len(parts):
                                return ManpuResult(False, f"Error: Arquivo xlsx tem células insuficientes")
                            ws.cell(row=row, column=int(column) + 1, value=part)
                else:
                    ws.cell(row=int(cell_val) + 1, column=int(column) + 1, value=text)

            else:
                # texto puro
                if cell_val == "full":
                    self._text = (self._text or "") + sep.join(parts)
                else:
                    self._text = (self._text or "") + text

            return ManpuResult(True)

        except Exception as e:
            return ManpuResult(False, f"Error: {e}")

    def read(self) -> Any:
        """Lê o conteúdo do arquivo aberto."""
        if self._ext == "csv" and self._rows is not None:
            if not self._rows:
                return []
            headers = self._rows[0] if self._rows else []
            return [dict(zip(headers, row)) for row in self._rows[1:]]
        elif self._ext in ("xlsx", "xls") and self._wb:
            ws = self._wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return []
            headers = [str(h) if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
            return [dict(zip(headers, row)) for row in rows[1:]]
        return self._text or ""

    def save(self) -> ManpuResult:
        """Salva as alterações no arquivo."""
        try:
            if self._ext == "csv" and self._rows is not None:
                with open(self._path, "w", newline="", encoding="utf-8") as f:
                    csv.writer(f).writerows(self._rows)
            elif self._ext in ("xlsx", "xls") and self._wb:
                self._wb.save(self._path)
            elif self._text is not None:
                self._path.write_text(self._text, encoding="utf-8")
            return ManpuResult(True)
        except Exception as e:
            return ManpuResult(False, f"Error: {e}")

    # Suporte ao `using ... as` da PoolScript
    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.save()
        if self._wb:
            self._wb.close()

    def __repr__(self):
        return f"<ManpuFile {self._path.name}>"


def open_file(target: str) -> ManpuFile:
    """
    Abre um arquivo pra operações — use com `using`:

        using mp.open(target="meu.xlsx") as arq {
            arq.write(column=0, cell=full, content=lista)
        }
    """
    return ManpuFile(str(target))


def src(filepath: str) -> str:
    """Resolve o caminho absoluto do arquivo."""
    return str(Path(filepath).resolve())


EXPORTS = {
    "read": read,
    "load": load,
    "write": write,
    "remove": remove,
    "open": open_file,
    "src": src,
}

